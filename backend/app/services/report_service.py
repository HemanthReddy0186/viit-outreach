import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from typing import List
from backend.app.models.models import DailyAttendance, Student, CallTask, ParentQuery

def generate_word_report_docx(
    date_str: str,
    section_name: str,
    coordinator_name: str,
    attendance_records: List[DailyAttendance],
    students: List[Student],
    calls: List[CallTask],
    queries: List[ParentQuery]
) -> io.BytesIO:
    doc = docx.Document()
    
    # Page setup: Standard Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title / Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("VIGNAN'S INSTITUTE OF INFORMATION TECHNOLOGY (AUTONOMOUS)\n")
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.color.rgb = RGBColor(15, 118, 110) # Forest Teal
    
    sub_title = title_p.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\nDAILY ATTENDANCE & PARENT VOICE OUTREACH REPORT")
    sub_title.bold = True
    sub_title.font.size = Pt(11)
    
    # Meta Details
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.rows[0].cells[0].text = f"Academic Date: {date_str}"
    meta_table.rows[0].cells[1].text = f"Section: CSE — Sec {section_name}"
    meta_table.rows[1].cells[0].text = f"Coordinator: {coordinator_name}"
    meta_table.rows[1].cells[1].text = f"Total Enrolled Strength: {len(students)}"
    
    doc.add_paragraph() # spacer
    
    # 1. Executive Attendance Summary
    h1 = doc.add_heading("1. Executive Attendance Summary", level=2)
    h1.runs[0].font.color.rgb = RGBColor(15, 118, 110)
    
    absentees = [a for a in attendance_records if a.status == "ABSENT"]
    present_count = len(students) - len(absentees)
    rate = ((present_count / len(students)) * 100) if students else 0
    
    sum_p = doc.add_paragraph()
    sum_p.add_run(f"• Total Students Enrolled: {len(students)}\n")
    sum_p.add_run(f"• Students Present Today: {present_count}\n")
    sum_p.add_run(f"• Absentees Requiring Parental Outreach: {len(absentees)}\n")
    sum_p.add_run(f"• Section Attendance Rate: {rate:.1f}%\n")
    
    # 2. Parental Voice Call Log
    h2 = doc.add_heading("2. Parental Voice Outreach Log (Prerecorded Telugu OBD / IVR)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(15, 118, 110)
    
    if calls:
        call_table = doc.add_table(rows=1, cols=5)
        call_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = call_table.rows[0].cells
        hdr_cells[0].text = "Roll Number"
        hdr_cells[1].text = "Student Name"
        hdr_cells[2].text = "Parent Phone"
        hdr_cells[3].text = "Call Status"
        hdr_cells[4].text = "Duration"
        
        for c in calls:
            row_cells = call_table.add_row().cells
            row_cells[0].text = c.roll_number
            row_cells[1].text = c.student_name
            row_cells[2].text = c.parent_phone
            row_cells[3].text = c.status
            row_cells[4].text = f"{c.duration_seconds}s"
    else:
        doc.add_paragraph("No parental voice calls recorded for this date.")
        
    doc.add_paragraph()
    
    # 3. Parent Queries & Follow-ups
    h3 = doc.add_heading("3. Parent Inquiries Routed to Coordinator", level=2)
    h3.runs[0].font.color.rgb = RGBColor(15, 118, 110)
    
    if queries:
        for i, q in enumerate(queries, 1):
            qp = doc.add_paragraph()
            qp.add_run(f"{i}. [{q.student_roll}] {q.student_name} — Category: {q.category}\n").bold = True
            qp.add_run(f"   Inquiry: \"{q.summary}\"\n")
            qp.add_run(f"   Status: {q.status} | Resolution: {q.resolution_notes or 'Pending review'}\n")
    else:
        doc.add_paragraph("No parent inquiries pending for this date.")
        
    doc.add_paragraph()
    
    # Sign-off block
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.rows[0].cells[0].text = "________________________\nDr. K. Srinivasa Rao\nClass Coordinator"
    sign_table.rows[0].cells[1].text = "________________________\nHead of Department\nDept of CSE, VIIT"
    
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
